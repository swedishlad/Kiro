#!/usr/bin/env python3
"""Reference parser: Mass.gov "Basic Service Rates" XLSX -> residential snapshot CSVs + rates.json.

Usage:
  pip install openpyxl==3.1.5
  python3 reference/ingest_rates.py data/raw/basic_service_rates_2026-07-02.xlsx data/build/
  # or, from the verified snapshot (no XLSX, no openpyxl needed):
  python3 reference/ingest_rates.py --from-snapshot data/snapshot/monthly_residential_cents_2026-07-02.csv \
          data/snapshot/fixed_residential_cents_2026-07-02.csv data/build/

Outputs (in out_dir):
  monthly_residential_cents.csv   year,month,NG,NSTAR,WMECo,Unitil   (newest first, same as workbook)
  fixed_residential_cents.csv     term_end_year,term_end_month,...    (value sits on the term's LAST month)
  rates.json                      contract in contracts/types.ts (RatesFile)
Then run reference/verify_snapshot.py on the two CSVs (gate G-DATA-1).

Layout facts this parser relies on (inspected Oct 22 2026; see docs/apis/01-massgov-basic-service-xlsx.md):
  sheets "Monthly Prices" and "Fixed Prices"; data from row 4 down, newest first;
  column A = year, only on the first row of each year block (carry it down);
  column B = month label (Jan..Dec, also "June", "July", "Sept");
  residential = columns D,E,F,G = National Grid, NSTAR (Eversource East), WMECo (Eversource West), Unitil;
  some numbers are stored as text with trailing "\r\n".
Tested against a replica workbook (reference/test_ingest.py), NOT against the live file in this sandbox.
"""
import csv
import json
import os
import sys
from datetime import date

MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
          "aug": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12}
UTIL_COLS = [("NG", 4), ("NSTAR", 5), ("WMECo", 6), ("Unitil", 7)]     # 1-based column indexes D..G
UTIL_IDS = {"NG": "NG", "NSTAR": "EV_EAST", "WMECo": "EV_WEST", "Unitil": "UNITIL"}


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 3)
    s = str(v).strip()
    if s == "":
        return None
    try:
        return round(float(s), 3)
    except ValueError:
        raise ValueError(f"non-numeric rate cell: {v!r}")


def read_sheet(ws):
    """Return [(year, month, [NG, NSTAR, WMECo, Unitil])] in workbook order (newest first)."""
    out, year = [], None
    for r in range(4, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, (int, float)):
            year = int(a)
        label = ws.cell(r, 2).value
        m = MONTHS.get(str(label).strip().lower()) if label is not None else None
        if m is None:
            continue
        vals = [num(ws.cell(r, c).value) for _, c in UTIL_COLS]
        if any(v is not None for v in vals):
            if year is None:
                raise ValueError(f"row {r}: month without a year above it")
            out.append((year, m, vals))
    return out


def ym(y, m):
    return f"{y:04d}-{m:02d}"


def add_months(y, m, k):
    t = y * 12 + (m - 1) + k
    return t // 12, t % 12 + 1


def derive_terms(fixed_rows):
    """Term start = previous term end + 1 month, per utility. First term's start is unknown."""
    terms = []
    for j, (name, _) in enumerate(UTIL_COLS):
        ends = sorted((y, m, vals[j]) for y, m, vals in fixed_rows if vals[j] is not None)
        prev = None
        for y, m, cents in ends:
            if prev is None:
                start, months = None, None
            else:
                sy, sm = add_months(prev[0], prev[1], 1)
                start = ym(sy, sm)
                months = (y * 12 + m) - (sy * 12 + sm) + 1
            terms.append({"utility": UTIL_IDS[name], "start": start, "end": ym(y, m),
                          "months": months, "cents": cents})
            prev = (y, m)
    return terms


def write_csv(path, header, rows):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        for y, m, vals in rows:
            w.writerow([y, m] + ["" if v is None else format(v, "g") for v in vals])


def read_csv(path):
    with open(path, newline="") as fh:
        rd = csv.reader(fh)
        next(rd)
        return [(int(r[0]), int(r[1]), [num(x) for x in r[2:6]]) for r in rd if any(x.strip() for x in r[2:6])]


def main(xlsx, out_dir, snapshot=None):
    if snapshot:
        monthly, fixed = read_csv(snapshot[0]), read_csv(snapshot[1])
        source_file = os.path.basename(snapshot[1]).replace("fixed_residential_cents_", "snapshot_")
    else:
        from openpyxl import load_workbook          # imported here so --help works without it
        wb = load_workbook(xlsx, data_only=True, read_only=False)
        for need in ("Monthly Prices", "Fixed Prices"):
            if need not in wb.sheetnames:
                raise SystemExit(f"missing sheet {need!r}; found {wb.sheetnames}")
        monthly = read_sheet(wb["Monthly Prices"])
        fixed = read_sheet(wb["Fixed Prices"])
        source_file = os.path.basename(xlsx)
    os.makedirs(out_dir, exist_ok=True)
    mpath = os.path.join(out_dir, "monthly_residential_cents.csv")
    fpath = os.path.join(out_dir, "fixed_residential_cents.csv")
    write_csv(mpath, ["year", "month", "NG", "NSTAR", "WMECo", "Unitil"], monthly)
    write_csv(fpath, ["term_end_year", "term_end_month", "NG", "NSTAR", "WMECo", "Unitil"], fixed)
    terms = derive_terms(fixed)
    current = {}
    for t in terms:
        if t["utility"] not in current or t["end"] > current[t["utility"]]["termEnd"]:
            current[t["utility"]] = {"termStart": t["start"], "termEnd": t["end"], "cents": t["cents"]}
    rates = {
        "source": {"id": "S1", "file": source_file, "retrieved": date.today().isoformat(),
                   "checksum": "run reference/verify_snapshot.py"},
        "monthly": [{"ym": ym(y, m), "utility": UTIL_IDS[n], "cents": vals[j]}
                    for y, m, vals in sorted(monthly) for j, (n, _) in enumerate(UTIL_COLS)
                    if vals[j] is not None],
        "fixedTerms": terms,
        "current": current,
    }
    with open(os.path.join(out_dir, "rates.json"), "w") as fh:
        json.dump(rates, fh, indent=1)
    odd = [t for t in terms if t["months"] not in (None, 6)]
    print(f"monthly rows={len(monthly)} fixed rows={len(fixed)} terms={len(terms)}")
    print("current:", json.dumps(current))
    print(f"non-6-month terms ({len(odd)}) -- verify each against the per-utility tabs:")
    for t in odd:
        print(f"  {t['utility']:8s} {t['start']} .. {t['end']}  months={t['months']}  {t['cents']}")


if __name__ == "__main__":
    if len(sys.argv) == 5 and sys.argv[1] == "--from-snapshot":
        main(None, sys.argv[4], snapshot=(sys.argv[2], sys.argv[3]))
    elif len(sys.argv) == 3:
        main(sys.argv[1], sys.argv[2])
    else:
        sys.exit(__doc__)
