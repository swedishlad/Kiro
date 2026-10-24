#!/usr/bin/env python3
"""Extract the Massachusetts usage data FairWatt Home needs from the two federal sources.

Usage:
  python3 reference/ingest_usage.py data/raw/recs2020_public_v7.csv data/raw/sales_revenue.xlsx data/build/
  python3 reference/ingest_usage.py --from-snapshot data/snapshot/recs_ma_2020.csv \
                                    data/snapshot/eia_ma_monthly_2026-08.csv data/build/
Writes data/build/recs_ma.csv, data/build/eia_ma_monthly.csv and data/build/eia_ma_total_monthly.csv (all sectors).
Verify with reference/verify_usage.py (G-DATA-3). The --from-snapshot form takes an optional 4th file: the
all-sector snapshot, copied to eia_ma_total_monthly.csv.

Sources (docs/apis/08-eia-usage-data.md):
  S14 RECS 2020 public microdata v7: one row per surveyed household; we keep the 552 rows with state_postal == "MA"
      and the columns in RECS_COLS, values copied as the source writes them (no re-formatting).
  S15 EIA-861M sales_revenue.xlsx, sheet "Monthly-States": rows with State == "MA"; residential columns E-H
      (revenue $k, sales MWh, customers, price c/kWh). kwh_per_customer = sales_mwh * 1000 / customers, 3 dp.
Needs openpyxl only for the XLSX path.
"""
import csv
import os
import shutil
import sys

RECS_COLS = ["DOEID", "KWH", "DOLLAREL", "TOTSQFT_EN", "YEARMADERANGE", "FUELHEAT", "EQUIPM", "TYPEHUQ",
             "NHSLDMEM", "HDD65", "CDD65", "NWEIGHT"]
EIA_COLS = ["year", "month", "status", "revenue_kusd", "sales_mwh", "customers", "price_cents", "kwh_per_customer"]
TOTAL_COLS = ["year", "month", "status", "total_revenue_kusd", "total_sales_mwh", "total_customers", "total_price_cents"]
SHEET = "Monthly-States"


def num(v):
    """Write a workbook number the same way on every run: ints stay ints, floats use Python's shortest repr."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def recs(src, out):
    with open(src, newline="") as fh, open(out, "w", newline="") as fo:
        rd = csv.DictReader(fh)
        missing = [c for c in RECS_COLS + ["state_postal"] if c not in rd.fieldnames]
        if missing:
            sys.exit(f"RECS file is missing columns {missing}")
        wr = csv.writer(fo, lineterminator="\n")
        wr.writerow(RECS_COLS)
        n = 0
        for r in rd:
            if r["state_postal"] == "MA":
                wr.writerow([r[c] for c in RECS_COLS])
                n += 1
    print(f"recs_ma: {n} rows -> {out}")


def eia(src, out):
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[SHEET]
    head = [c for c in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
    if head[:4] != ["Year", "Month", "State", "Data Status"]:
        sys.exit(f"unexpected header row 3 in {SHEET}: {head[:8]}")
    rows = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r or r[2] != "MA":
            continue
        year, month, status, rev, sales, cust, price = int(r[0]), int(r[1]), r[3], r[4], r[5], r[6], r[7]
        if not sales or not cust:
            sys.exit(f"MA {year}-{month}: missing residential sales or customers")
        rows.append([year, month, status, num(rev), num(sales), num(cust), num(price),
                     f"{sales * 1000 / cust:.3f}"])
    rows.sort(key=lambda x: (x[0], x[1]))
    with open(out, "w", newline="") as fo:
        wr = csv.writer(fo, lineterminator="\n")
        wr.writerow(EIA_COLS)
        wr.writerows(rows)
    print(f"eia_ma_monthly: {len(rows)} rows {rows[0][0]}-{rows[0][1]:02d} .. {rows[-1][0]}-{rows[-1][1]:02d} -> {out}")


def eia_total(src, out):
    """All sectors (columns U-X, block 'TOTAL'): the level reference for ISO-NE load after its data ends (2023)."""
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[SHEET]
    sectors = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    if sectors[20] != "TOTAL":
        sys.exit(f"expected the TOTAL block at column U in {SHEET}, got {sectors[20]!r}")
    rows = [[int(r[0]), int(r[1]), r[3], num(r[20]), num(r[21]), num(r[22]), num(r[23])]
            for r in ws.iter_rows(min_row=4, values_only=True) if r and r[2] == "MA"]
    rows.sort(key=lambda x: (x[0], x[1]))
    with open(out, "w", newline="") as fo:
        wr = csv.writer(fo, lineterminator="\n")
        wr.writerow(TOTAL_COLS)
        wr.writerows(rows)
    print(f"eia_ma_total_monthly: {len(rows)} rows -> {out}")


def main():
    args = sys.argv[1:]
    if len(args) in (4, 5) and args[0] == "--from-snapshot":
        rsrc, esrc, *tsrc, outdir = args[1:]
        os.makedirs(outdir, exist_ok=True)
        shutil.copyfile(rsrc, os.path.join(outdir, "recs_ma.csv"))
        shutil.copyfile(esrc, os.path.join(outdir, "eia_ma_monthly.csv"))
        if tsrc:
            shutil.copyfile(tsrc[0], os.path.join(outdir, "eia_ma_total_monthly.csv"))
        print(f"copied snapshot -> {outdir}/ ({2 + len(tsrc)} files)")
        return
    if len(args) != 3:
        sys.exit(__doc__)
    rsrc, esrc, outdir = args
    os.makedirs(outdir, exist_ok=True)
    recs(rsrc, os.path.join(outdir, "recs_ma.csv"))
    eia(esrc, os.path.join(outdir, "eia_ma_monthly.csv"))
    eia_total(esrc, os.path.join(outdir, "eia_ma_total_monthly.csv"))


if __name__ == "__main__":
    main()
