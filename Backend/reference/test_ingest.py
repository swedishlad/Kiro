#!/usr/bin/env python3
"""Test ingest_rates.py on a REPLICA workbook built from the verified snapshot CSVs.

It reproduces the quirks seen in the real file: header rows 1-2, data from row 4, newest first,
year only on the first row of each year block (starting at Dec), long month labels (June/July/Sept),
empty future rows, a "market" text cell outside the residential columns, and one number stored
as text with "\r\n". Passing this test proves the parser logic, not that the live file is unchanged.

Usage: python3 reference/test_ingest.py   (needs openpyxl==3.1.5)
"""
import csv
import filecmp
import os
import subprocess
import sys
import tempfile

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SNAP = os.path.join(HERE, "..", "data", "snapshot")
LABELS = {1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "May", 6: "June", 7: "July", 8: "Aug",
          9: "Sept", 10: "Oct", 11: "Nov", 12: "Dec"}


def load(path):
    with open(path) as fh:
        rd = csv.reader(fh)
        next(rd)
        return {(int(r[0]), int(r[1])): r[2:6] for r in rd}


def build_sheet(ws, data, text_cell=None):
    ws.cell(1, 4, "Residential")
    for c, n in zip(range(4, 8), ["National Grid", "NSTAR", "WMECo", "Unitil"]):
        ws.cell(2, c, n)
    row = 4
    for y in range(2027, 2004, -1):
        for m in range(12, 0, -1):
            if m == 12:
                ws.cell(row, 1, y)
            ws.cell(row, 2, LABELS[m])
            ws.cell(row, 22, "market")                      # junk text outside residential columns
            vals = data.get((y, m))
            if vals:
                for j, v in enumerate(vals):
                    if v != "":
                        cell_val = float(v)
                        if text_cell == (y, m, j):
                            cell_val = v + "\r\n"           # number stored as text
                        ws.cell(row, 4 + j, cell_val)
            row += 1


def main():
    monthly = load(os.path.join(SNAP, "monthly_residential_cents_2026-07-02.csv"))
    fixed = load(os.path.join(SNAP, "fixed_residential_cents_2026-07-02.csv"))
    with tempfile.TemporaryDirectory() as tmp:
        wb = Workbook()
        wb.active.title = "Overview"
        build_sheet(wb.create_sheet("Monthly Prices"), monthly)
        build_sheet(wb.create_sheet("Fixed Prices"), fixed, text_cell=(2015, 10, 0))
        xlsx = os.path.join(tmp, "replica.xlsx")
        wb.save(xlsx)
        out = os.path.join(tmp, "out")
        subprocess.run([sys.executable, os.path.join(HERE, "ingest_rates.py"), xlsx, out], check=True)
        res = subprocess.run([sys.executable, os.path.join(HERE, "verify_snapshot.py"),
                              os.path.join(out, "monthly_residential_cents.csv"),
                              os.path.join(out, "fixed_residential_cents.csv")])
        same_m = filecmp.cmp(os.path.join(out, "monthly_residential_cents.csv"),
                             os.path.join(SNAP, "monthly_residential_cents_2026-07-02.csv"), shallow=False)
        same_f = filecmp.cmp(os.path.join(out, "fixed_residential_cents.csv"),
                             os.path.join(SNAP, "fixed_residential_cents_2026-07-02.csv"), shallow=False)
        print("byte-identical to snapshot:", same_m, same_f)
        ok = res.returncode == 0 and same_m and same_f
        print("TEST:", "PASS" if ok else "FAIL")
        sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
